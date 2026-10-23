from pathlib import Path
import json, csv
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Pt as WordPt
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
root=Path(__file__).resolve().parent
output=root/'exports'
output.mkdir(exist_ok=True)
slides=[
 {'kind':'cover','kicker':'Internal launch briefing','title':'A controlled shift to November 19.','subtitle':'Protect the retail launch buffer. Keep the November 12 press embargo.','items':[],'notes':'This is an internal draft for review. The recommendation does not authorize changing dates or sending external notices.'},
 {'kind':'timeline','kicker':'The sequence','title':'Three dates. One clear sequence.','subtitle':'Packaging arrives before the press embargo and retail availability.','timeline':[{'date':'Nov 8','label':'Packaging arrives'},{'date':'Nov 12','label':'Press embargo'},{'date':'Nov 19','label':'Retail availability'}],'items':[],'notes':'Packaging moved from November 4 to November 8. The extra week protects fulfilment. The press announcement should state the retail availability date clearly.'},
 {'kind':'list','kicker':'Impact','title':'What changes, what holds.','subtitle':'Keep the recommendation separate from actions that still need approval.','items':['Retail availability moves seven days to November 19.','The press embargo remains November 12.','Paid media requires its own review after source access is restored.'],'notes':'The media buy tracker was excluded from this recommendation because its access expired. The dates use the permitted supplier ETA, retail calendar, and plan snapshot.'},
 {'kind':'list','kicker':'Decisions','title':'Two approvals. Two boundaries.','subtitle':'A date decision never silently authorizes external delivery.','items':['Jamie Davis approves or redirects the proposed retail date.','Four vendor notices are reviewed separately before sending.','Ops holds delivery until that specific action is approved.'],'notes':'These are proposed actions in a sample design. No notices were sent. Each approval should identify its actor, target, and exact content or change.'},
 {'kind':'list','kicker':'Ownership','title':'The next step has an owner.','subtitle':'Keep progress visible without asking people to supervise every tool call.','items':['Launch Planner maintains the plan and completion criteria.','Ops prepares vendor notices and holds delivery.','Jon T. restores and verifies access to the media source.'],'notes':'A question about presentation audience blocks only the briefing assignment. Date approval, source restoration, and message approval remain separate work.'},
 {'kind':'list','kicker':'Evidence and follow-through','title':'Grounded, versioned, and ready to review.','subtitle':'Every output points back to the work and the evidence that produced it.','items':['Supplier ETA: indexed October 23 at 09:42 UTC.','Retail calendar: live search as Jamie at 09:47 UTC.','Launch plan v3: snapshot imported October 18.'],'notes':'This briefing is version 1, a draft. Review the outcome, request slide-specific changes, and retain prior versions when new outputs are produced.'}
]
team=json.loads(json.dumps(slides))
team[0].update(title='The launch plan, with a clear next step.',subtitle='A project-team briefing for the November 19 retail launch.')
team[2]['title']='What the team needs to adjust.'
team[4]['title']='Owners, blockers, and next actions.'
decks={'leadership':{'audience':'Leadership','slides':slides},'project-team':{'audience':'Project team','slides':team}}
(root/'presentation-data.json').write_text(json.dumps(decks,ensure_ascii=False,indent=2))
ink='1B2938';accent='2E5B88';muted='52687B'
def box(slide,x,y,w,h,text,size=20,color=ink,bold=False):
    frame=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    frame.word_wrap=True
    frame.margin_left=frame.margin_right=0
    frame.margin_top=frame.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
        p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(12)
    return frame
for key,deck in decks.items():
    pres=Presentation();pres.slide_width=Inches(13.333);pres.slide_height=Inches(7.5)
    for number,s in enumerate(deck['slides'],1):
        slide=pres.slides.add_slide(pres.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('EAF0F6' if s['kind']=='cover' else 'FDFDFC')
        for r in range(2):
            for col in range(2):
                shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.65+col*.14),Inches(.42+r*.14),Inches(.10),Inches(.10))
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(accent);shape.line.fill.background();shape.rotation=352
        box(slide,1.04,.40,4,.35,'humagen / Spring Launch',12,accent,True)
        box(slide,.7,1.24,11.8,.3,s['kicker'].upper(),12,muted,True)
        box(slide,.7,1.86,11.8,1.28,s['title'],36,ink,True)
        box(slide,.7,3.13,11.8,.65,s['subtitle'],20,muted)
        if s.get('timeline'):
            for j,m in enumerate(s['timeline']):
                x=.7+j*4.05
                rule=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(4.2),Inches(3.6),Inches(.025));rule.fill.solid();rule.fill.fore_color.rgb=RGBColor.from_string(accent);rule.line.fill.background()
                box(slide,x,4.5,3.6,.55,m['date'],32,ink,True);box(slide,x,5.2,3.6,.5,m['label'],18,muted)
        else:
            for j,item in enumerate(s['items']):box(slide,.85,4.08+j*.62,11.4,.6,'• '+item,21,ink)
        box(slide,.7,7.0,10,.25,'INTERNAL DRAFT / SAMPLE DATA / '+deck['audience'].upper(),10,muted)
        box(slide,12.05,7.0,.6,.25,f'{number} / 6',10,muted)
        slide.notes_slide.notes_text_frame.text=s['notes']
    pres.save(output/f'Launch-briefing-{key}.pptx')
fontroot=Path('/usr/share/fonts/truetype/dejavu')
pdfmetrics.registerFont(TTFont('HG',str(fontroot/'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('HGBold',str(fontroot/'DejaVuSans-Bold.ttf')))
def para(c,value,x,y,w,size=18,bold=False,color=ink):
    st=ParagraphStyle('p',fontName='HGBold' if bold else 'HG',fontSize=size,leading=size*1.3,textColor=colors.HexColor('#'+color))
    p=Paragraph(value,st);_,h=p.wrap(w,900);p.drawOn(c,x,y-h);return y-h
def mark(c,x,y,size=18):
    c.saveState();c.translate(x+size/2,y+size/2);c.rotate(8)
    c.setFillColor(colors.HexColor('#'+accent))
    for row in range(2):
        for col in range(2):c.roundRect(-size/2+col*size*.57,-size/2+row*size*.57,size*.40,size*.40,size*.10,fill=1,stroke=0)
    c.restoreState()
for key,deck in decks.items():
    c=Canvas(str(output/f'Launch-briefing-{key}.pdf'),pagesize=(960,540));c.setTitle('Launch briefing - '+deck['audience'])
    for number,s in enumerate(deck['slides'],1):
        c.setFillColor(colors.HexColor('#EAF0F6' if s['kind']=='cover' else '#FDFDFC'));c.rect(0,0,960,540,fill=1,stroke=0)
        mark(c,50,489,18);para(c,'humagen / Spring Launch',77,505,793,11,True,accent)
        para(c,s['kicker'].upper(),50,435,840,11,True,muted)
        para(c,s['title'],50,385,850,32,True)
        para(c,s['subtitle'],50,288,850,17,False,muted)
        if s.get('timeline'):
            for j,m in enumerate(s['timeline']):
                x=50+j*295;c.setStrokeColor(colors.HexColor('#'+accent));c.setLineWidth(2);c.line(x,212,x+260,212)
                para(c,m['date'],x,190,260,31,True);para(c,m['label'],x,138,260,16,False,muted)
        else:
            for j,item in enumerate(s['items']):para(c,'• '+item,60,215-j*46,830,17)
        para(c,'INTERNAL DRAFT / SAMPLE DATA / '+deck['audience'].upper(),50,30,750,9,False,muted);para(c,f'{number} / 6',870,30,70,9,False,muted);c.showPage()
    c.save()
plan=[('Recommendation','Move retail availability to November 19. Keep the press embargo on November 12.'),('Why','Packaging now arrives November 8, four days later than planned. The revised retail date restores the fulfilment buffer.'),('Milestones','November 8: packaging arrives. November 12: press embargo. November 19: retail availability.'),('Approvals','Jamie Davis reviews the date. Four vendor notices require a separate delivery approval.'),('Owners','Launch Planner owns the plan. Ops prepares notices. Jon T. restores access to the media buy tracker.'),('Evidence','Supplier ETA indexed October 23, 09:42 UTC. Retail calendar queried as Jamie at 09:47 UTC. Launch plan v3 imported October 18.'),('Status','Internal draft for review. No external messages have been sent from this sample.')]
doc=Document();doc.add_heading('Launch plan',0);doc.add_paragraph('Version 4 · Spring Launch · Internal draft · Sample data')
for heading,body in plan:doc.add_heading(heading,level=1);doc.add_paragraph(body)
doc.core_properties.title='Launch plan v4';doc.core_properties.author='Humagen / Launch Planner';doc.save(output/'Launch-plan-v4.docx')
c=Canvas(str(output/'Launch-plan-v4.pdf'),pagesize=(612,792));c.setTitle('Launch plan v4');y=738
mark(c,48,y-16,16);y=para(c,'humagen / Spring Launch',73,y,491,10,True,accent)-24
y=para(c,'Launch plan',48,y,516,28,True)-12
y=para(c,'Version 4 · Internal draft · Sample data',48,y,516,10,False,muted)-24
for heading,body in plan:
    y=para(c,heading,48,y,516,12,True)-6;y=para(c,body,48,y,516,11)-20
c.save()
rows=[['Dependency','Owner','Current date','Status'],['Packaging','Jon T.','Nov 8','Delayed four days'],['Press embargo','Mira K.','Nov 12','Holds'],['Retail availability','Jamie Davis','Nov 19','Needs approval'],['Vendor notices','Ops','After approval','Four drafts held'],['Paid media','Sam R.','Pending review','Source access blocked']]
with (output/'Launch-dependencies.csv').open('w',newline='') as f:csv.writer(f).writerows(rows)
(root/'dependencies.json').write_text(json.dumps(rows))
(root/'Humagen-mark.svg').write_text('''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 32 32"><title>Humagen mark</title><g fill="#1b1c1e" transform="rotate(-8 16 16)"><rect x="5" y="5" width="9" height="9" rx="2.5"/><rect x="18" y="5" width="9" height="9" rx="2.5"/><rect x="5" y="18" width="9" height="9" rx="2.5"/><rect x="18" y="18" width="9" height="9" rx="2.5"/></g></svg>''')
print('Generated 2 editable presentation variants, their PDFs, plan DOCX/PDF, dependency CSV, and logo SVG.')
